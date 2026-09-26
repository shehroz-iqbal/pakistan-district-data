"""Step 4: write the published files.

    python -m pipeline.export

data/processed/  CSVs and an Excel workbook for people who want the numbers
site/data/       the compact JSON + GeoJSON the dashboard reads
"""

from __future__ import annotations

import json
import shutil
import sys

import duckdb
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .common import DB_PATH, PROCESSED, SITE_DATA, load_yaml, sources_pinned

WORKBOOK = PROCESSED / "pakistan_district_indicators_2023.xlsx"


def frames(con) -> dict[str, pd.DataFrame]:
    catalog = con.execute("SELECT * FROM indicator_catalog ORDER BY sort_order").df()
    order = list(catalog["indicator_id"])

    districts = con.execute("""
        SELECT district_id, district_name, province_code, province_name, census_name, map_unit_id
        FROM dim_district ORDER BY province_code, district_name""").df()

    wide = con.execute("SELECT * FROM district_indicators").df()
    wide = wide[["district_id", "district_name", "province_code", "province_name"] + order]
    for _, ind in catalog.iterrows():
        wide[ind["indicator_id"]] = wide[ind["indicator_id"]].round(int(ind["decimals"]))

    long = con.execute("""
        WITH names AS (
            SELECT district_id AS geo_id, district_name AS geo_name FROM dim_district
            UNION ALL SELECT DISTINCT map_unit_id, map_unit_name FROM dim_district
                      WHERE map_unit_id <> district_id
            UNION ALL SELECT DISTINCT province_code, province_name FROM dim_district
            UNION ALL SELECT 'PK', 'Pakistan'
        )
        SELECT v.indicator_id, v.level, v.geo_id, n.geo_name,
               round(v.value, 6) AS value, v.numerator, v.denominator, v.n_districts
        FROM indicator_values v JOIN names n USING (geo_id)
        JOIN indicator_catalog c USING (indicator_id)
        ORDER BY c.sort_order,
                 CASE v.level WHEN 'national' THEN 0 WHEN 'province' THEN 1
                              WHEN 'map_unit' THEN 2 ELSE 3 END,
                 v.geo_id""").df()

    counts = con.execute("""
        SELECT * EXCLUDE (district_name, census_name, province_name, map_unit_id, map_unit_name)
        FROM district_counts ORDER BY province_code, district_id""").df()

    public_catalog = catalog[["indicator_id", "label", "theme", "unit", "decimals", "kind",
                              "numerator_col", "denominator_col", "multiplier", "direction",
                              "source_table", "definition"]]
    return {"catalog": public_catalog, "districts": districts, "wide": wide,
            "long": long, "counts": counts}


def write_csvs(f: dict[str, pd.DataFrame]) -> None:
    f["districts"].to_csv(PROCESSED / "districts.csv", index=False)
    f["wide"].to_csv(PROCESSED / "district_indicators.csv", index=False)
    f["long"].to_csv(PROCESSED / "indicators_long.csv", index=False)
    f["catalog"].to_csv(PROCESSED / "indicator_catalog.csv", index=False)
    f["counts"].to_csv(PROCESSED / "district_counts.csv", index=False)


def write_workbook(f: dict[str, pd.DataFrame]) -> None:
    """A workbook for colleagues who live in Excel: labelled columns, a notes sheet."""
    catalog = f["catalog"].set_index("indicator_id")
    labels = {i: f"{catalog.at[i, 'label']} ({catalog.at[i, 'unit']})" for i in catalog.index}

    districts = f["wide"].rename(columns={"district_name": "District", "province_name": "Province",
                                          "district_id": "ID", **labels})
    districts = districts.drop(columns=["province_code"])

    summary = f["long"][f["long"]["level"].isin(["national", "province"])]
    summary = summary.pivot_table(index="geo_name", columns="indicator_id", values="value", sort=False)
    summary = summary[list(catalog.index)]
    for i in catalog.index:
        summary[i] = summary[i].round(int(catalog.at[i, "decimals"]))
    summary = summary.rename(columns=labels).reset_index().rename(columns={"geo_name": "Area"})

    definitions = f["catalog"][["label", "unit", "theme", "definition", "source_table", "indicator_id"]]
    definitions = definitions.rename(columns=str.capitalize).rename(
        columns={"Source_table": "Census table", "Indicator_id": "Column ID"})

    readme = pd.DataFrame({"Pakistan district indicators, Census 2023": [
        f"Built by the pakistan-district-data pipeline from sources pinned on {sources_pinned():%d %B %Y}.",
        "Source: Pakistan Bureau of Statistics, 7th Population and Housing Census 2023 (district tables),",
        "via the CSV conversion by Fahad Mirza (MIT licence), with corrections listed in the repository.",
        "Province and Pakistan rows are rebuilt from district counts; they match PBS's published figures.",
        "Sheets: Districts (136 rows), Provinces (Pakistan + 5), Definitions (how each column is computed).",
        "Percentages are shares of the population or households named in each definition.",
    ]})

    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as xl:
        readme.to_excel(xl, sheet_name="Read me", index=False)
        districts.to_excel(xl, sheet_name="Districts", index=False)
        summary.to_excel(xl, sheet_name="Provinces", index=False)
        definitions.to_excel(xl, sheet_name="Definitions", index=False)
        header = PatternFill("solid", fgColor="1F4E5F")
        for ws in xl.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "B2" if ws.title in ("Districts", "Provinces") else "A2"
            for idx, col in enumerate(ws.columns, start=1):
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(idx)].width = min(max(10, width * 0.9), 60)
            ws.row_dimensions[1].height = 45 if ws.title in ("Districts", "Provinces") else 18


def write_site_data(f: dict[str, pd.DataFrame], con) -> None:
    """One JSON file with everything the dashboard needs, rounded for display."""
    SITE_DATA.mkdir(parents=True, exist_ok=True)
    catalog = f["catalog"]
    decimals = dict(zip(catalog["indicator_id"], catalog["decimals"], strict=True))
    long = f["long"]

    places: dict[str, dict] = {}
    for row in long.itertuples():
        place = places.setdefault(row.geo_id, {"name": row.geo_name, "level": row.level, "values": {}})
        if pd.notna(row.value):
            # Keep three spare decimals: the page rounds for display, and rounding
            # twice (60.652 -> 60.65 -> 60.6) would show the wrong last digit.
            place["values"][row.indicator_id] = round(float(row.value), int(decimals[row.indicator_id]) + 3)

    for d in f["districts"].itertuples():
        places[d.district_id].update(province=d.province_code, map_unit=d.map_unit_id)
    for mu, members in f["districts"].groupby("map_unit_id")["district_id"]:
        if mu in places and places[mu]["level"] == "map_unit":
            places[mu].update(province=f["districts"].set_index("map_unit_id").loc[mu, "province_code"].iloc[0],
                              districts=list(members))

    qa_path = PROCESSED / "QA.md"
    sources = load_yaml("sources.yml")
    payload = {
        "sources_pinned": sources_pinned().isoformat(),
        "indicators": catalog[["indicator_id", "label", "theme", "unit", "decimals",
                               "direction", "source_table", "definition"]].to_dict("records"),
        "provinces": dict(con.execute(
            "SELECT DISTINCT province_code, province_name FROM dim_district ORDER BY 1").fetchall()),
        "places": places,
        "sources": {
            "census": {"publisher": sources["census_2023"]["publisher"],
                       "url": sources["census_2023"]["publisher_url"],
                       "conversion": sources["census_2023"]["conversion"]["repo"]},
            "boundaries": {"publisher": sources["geoboundaries_pak"]["publisher"],
                           "url": sources["geoboundaries_pak"]["publisher_url"]},
        },
        "qa_summary": qa_path.read_text(encoding="utf-8").split("**")[1] if qa_path.exists() else None,
    }
    with open(SITE_DATA / "dashboard.json", "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, separators=(",", ":"))
    shutil.copyfile(PROCESSED / "districts_2023.geojson", SITE_DATA / "districts_2023.geojson")
    shutil.copyfile(PROCESSED / "provinces_2023.geojson", SITE_DATA / "provinces_2023.geojson")

    # Downloads offered on the dashboard itself.
    downloads = SITE_DATA / "downloads"
    downloads.mkdir(exist_ok=True)
    for name in ("district_indicators.csv", "indicator_catalog.csv", WORKBOOK.name):
        shutil.copyfile(PROCESSED / name, downloads / name)


def write_data_dictionary(f: dict[str, pd.DataFrame]) -> None:
    """docs/data_dictionary.md, generated so it can never disagree with the files."""
    cat = f["catalog"]
    lines = [
        "# Data dictionary",
        "",
        "Generated by `pipeline/export.py` from `config/indicators.yml`; edit that file, not this one.",
        "All files are in `data/processed/` (UTF-8 CSV, comma-separated, one header row).",
        "",
        "## `districts.csv`: one row per 2023 census district (136)",
        "",
        "| Column | Meaning |",
        "|---|---|",
        "| `district_id` | Stable identifier: ISO 3166-2 province code + PBS name, e.g. `pb_dera_ghazi_khan`. Never changes once published. |",
        "| `district_name` | Display name (PBS spelling, title case). |",
        "| `province_code`, `province_name` | ISO 3166-2:PK code (`PB`, `SD`, `KP`, `BA`, `IS`) and name. |",
        "| `census_name` | Name exactly as PBS prints it, for matching back to the tables. |",
        "| `map_unit_id` | The shape the district is drawn as on the map: itself, or `sd_karachi` for Karachi's seven districts. |",
        "",
        "## `district_indicators.csv`: one row per district, one column per indicator",
        "",
        "Values are rounded to the decimals shown below. Unrounded values, numerators and",
        "denominators are in `indicators_long.csv`.",
        "",
        "| Column | Label | Unit | Definition | Census table |",
        "|---|---|---|---|---|",
    ]
    for row in cat.itertuples():
        lines.append(f"| `{row.indicator_id}` | {row.label} | {row.unit} | {row.definition} | {row.source_table} |")
    lines += [
        "",
        "## `indicators_long.csv`: every indicator at every level",
        "",
        "| Column | Meaning |",
        "|---|---|",
        "| `indicator_id` | See the table above. |",
        "| `level` | `national`, `province`, `map_unit` (Karachi combined) or `district`. |",
        "| `geo_id`, `geo_name` | `PK`, a province code, `sd_karachi`, or a `district_id`. |",
        "| `value` | The indicator, unrounded (6 decimals). |",
        "| `numerator`, `denominator` | Summed counts behind a ratio (empty for differences). Province and national values are built from these, never by averaging district rates. |",
        "| `n_districts` | How many districts contributed. |",
        "",
        "## `district_counts.csv`: the counts every indicator is built from",
        "",
        "One row per district. Prefixes: `pop_` persons, `hh_` households, `owner_` dwelling owners;",
        "suffixes `_all`, `_male`, `_female`. `households_t20`, `_t23`, `_t24`, `_t25` are each table's own",
        "household total (the checks confirm they are identical). `pbs_` columns are figures PBS itself",
        "publishes (growth rate, density, ...), kept to cross-check our recomputations. Exact derivations:",
        "`sql/02_district_counts.sql`.",
        "",
        "## `districts_2023.geojson` and `provinces_2023.geojson`",
        "",
        "WGS84, RFC 7946 winding, coordinates rounded to 0.0001° (about 10 m) and simplified for web use.",
        "District features carry `map_unit_id`, `name`, `province_code`, `district_ids` (the census districts",
        "the shape covers) and `in_census_tables` (false for Azad Jammu & Kashmir and Gilgit-Baltistan).",
        "",
        "## `qa_district_area_check.csv`",
        "",
        "For each map shape: the area PBS reports, the area of the rebuilt shape, their ratio, and how its",
        "tehsil polygons were assigned (see `docs/methodology.md`).",
        "",
    ]
    (PROCESSED.parents[1] / "docs" / "data_dictionary.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    con = duckdb.connect(str(DB_PATH), read_only=True)
    f = frames(con)
    write_csvs(f)
    write_workbook(f)
    write_site_data(f, con)
    write_data_dictionary(f)
    con.close()
    size = (SITE_DATA / "dashboard.json").stat().st_size / 1e3
    print(f"export: {len(f['wide'])} districts x {len(f['catalog'])} indicators -> data/processed/ "
          f"(+ workbook); site/data/dashboard.json {size:.0f} kB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
